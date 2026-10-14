"""Pink Sheet monthly observations with explicit workbook/unit selection.

Official landing page: https://www.worldbank.org/en/research/commodity-markets .
The linked workbook is a changing latest file. Its month labels are economic
periods; receipt or an explicit file-version publication controls availability.
Data rights follow dataset/provider metadata; generic CC BY is not universal.
"""

import hashlib
import io
import re
import zipfile
from dataclasses import asdict, dataclass
from urllib.parse import urlsplit

import pandas as pd

from pfe.data.contracts import _instant, observations_asof
from pfe.data.eia import (
    ContextError,
    ContextLimitError,
    _collect_file,
    _identifier,
    _observation,
    _observations,
    _period,
)

PINK_SHEET_URL = (
    "https://thedocs.worldbank.org/en/doc/"
    "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
)
PINK_SHEET_LANDING_URL = "https://www.worldbank.org/en/research/commodity-markets"


@dataclass(frozen=True)
class CommodityColumn:
    entity_id: str
    unit: str
    metric: str = "price"
    currency: str | None = None

    def __post_init__(self):
        for name in ("entity_id", "unit", "metric"):
            _identifier(getattr(self, name), name)
        if self.currency is not None:
            _identifier(self.currency, "currency")


def _columns(column_specs):
    if not isinstance(column_specs, dict) or not column_specs:
        raise ContextError("Pink Sheet requires explicit column/entity/unit mapping")
    result = {}
    for name, spec in column_specs.items():
        if not isinstance(name, str) or not name:
            raise ContextError("Pink Sheet column names must be explicit nonempty strings")
        result[name] = CommodityColumn(**spec) if isinstance(spec, dict) else spec
        if not isinstance(result[name], CommodityColumn):
            raise ContextError("Pink Sheet columns require CommodityColumn values")
    return result


def parse_pink_sheet(
    frame,
    *,
    period_column,
    column_specs,
    received_at,
    revision_id,
    period_format="YYYY-MM",
    published_at=None,
    availability_basis="receipt",
    max_rows=25_000,
):
    """Convert a caller-selected monthly table, preserving original units.

    period_format accepts only explicit YYYY-MM or YYYYMmm declarations. Unit,
    currency, entity and value type come from column_specs, not header guesses.
    Values are not forward-filled or backdated to month end. publication mode
    needs an aware timestamp for this particular file version, not a report date.
    """
    _instant(received_at, "received_at")
    _identifier(revision_id, "revision_id")
    _identifier(period_column, "period_column")
    if availability_basis not in ("receipt", "publication"):
        raise ContextError("Pink Sheet availability must be receipt or publication")
    specs = _columns(column_specs)
    if not isinstance(frame, pd.DataFrame) or not frame.columns.is_unique:
        raise ContextError("Pink Sheet requires a DataFrame with unique columns")
    if not {period_column, *specs}.issubset(frame.columns):
        raise ContextError("Pink Sheet misses exact selected period/value columns")
    if type(max_rows) is not int or max_rows < 1 or len(frame) * len(specs) > max_rows:
        raise ContextLimitError("Pink Sheet selected observations exceed max_rows")
    if period_format not in ("YYYY-MM", "YYYYMmm"):
        raise ContextError("Pink Sheet period format must be explicit YYYY-MM or YYYYMmm")
    records = []
    for _, row in frame.iterrows():
        raw = row[period_column]
        if not isinstance(raw, str):
            raise ContextError("Pink Sheet economic periods require explicit monthly labels")
        if period_format == "YYYYMmm":
            if not re.fullmatch(r"\d{4}M\d{2}", raw):
                raise ContextError("Pink Sheet monthly label does not match YYYYMmm")
            period = raw[:4] + "-" + raw[5:]
        else:
            period = raw
        first, last, _ = _period(period, "monthly")
        for column, spec in specs.items():
            records.append(
                _observation(
                    entity_id=spec.entity_id,
                    metric=spec.metric,
                    period_start=first,
                    period_end=last,
                    period_label=raw,
                    value=row[column],
                    unit=spec.unit,
                    currency=spec.currency,
                    source_id="world_bank",
                    revision_id=revision_id,
                    received_at=received_at,
                    published_at=published_at,
                    availability_basis=availability_basis,
                    frequency="monthly",
                    calendar="Gregorian_months",
                )
            )
    return _observations(records)


def parse_pink_sheet_xlsx(
    body,
    *,
    sheet_name,
    header_row,
    period_column,
    column_specs,
    received_at,
    period_format="YYYY-MM",
    published_at=None,
    availability_basis="receipt",
    revision_id=None,
    max_rows=25_000,
    max_expanded_bytes=52_428_800,
    max_columns=256,
    data_start_row=None,
    data_end_row=None,
):
    """Read an explicit sheet/header without evaluating formulas or unbounded ZIPs.

    header_row is zero-based. openpyxl is optional (commodities extra); the core
    table parser has no Excel dependency. Workbook titles, units and blank rows
    are not guessed. A changed layout fails rather than selecting another sheet.
    """
    specs = _columns(column_specs)
    if type(max_rows) is not int or max_rows < 1:
        raise ContextError("Workbook max_rows must be positive")
    for name, value in (
        ("header_row", header_row),
        ("max_expanded_bytes", max_expanded_bytes),
        ("max_columns", max_columns),
    ):
        if type(value) is not int or value < (0 if name == "header_row" else 1):
            raise ContextError("Workbook positions/budgets require non-negative integers")
    if header_row > max_rows or not isinstance(body, bytes):
        raise ContextLimitError("Workbook header or bytes exceed the declared budget")
    first_data = header_row + 1 if data_start_row is None else data_start_row
    if type(first_data) is not int or first_data <= header_row:
        raise ContextError("data_start_row must explicitly follow the zero-based header row")
    if data_end_row is not None and (type(data_end_row) is not int or data_end_row < first_data):
        raise ContextError("data_end_row is an exclusive zero-based bound after data_start_row")
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if sum(item.file_size for item in archive.infolist()) > max_expanded_bytes:
                raise ContextLimitError("Pink Sheet expanded ZIP exceeds max_expanded_bytes")
    except zipfile.BadZipFile:
        raise ContextError("Pink Sheet requires an XLSX ZIP workbook") from None
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ContextError("Install the commodities extra to read Pink Sheet XLSX") from None
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ContextError("Explicit Pink Sheet sheet_name is absent")
        sheet = book[sheet_name]
        if sheet.max_column > max_columns:
            raise ContextLimitError("Pink Sheet worksheet exceeds max_columns")
        header = next(
            sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True)
        )
        if any(header.count(name) != 1 for name in specs):
            raise ContextError("Selected Pink Sheet column is absent or ambiguous")
        if type(period_column) is int:
            if not 0 <= period_column < len(header):
                raise ContextError("Explicit period column index falls outside worksheet header")
            period_position, period_name = period_column, "__economic_period__"
        else:
            if not isinstance(period_column, str) or header.count(period_column) != 1:
                raise ContextError("Pink Sheet period header is absent or ambiguous")
            period_position, period_name = header.index(period_column), period_column
        selected = [period_name, *specs]
        if period_position in [header.index(name) for name in specs]:
            raise ContextError("Period and selected commodity columns must be distinct")
        if not set(specs).issubset(header):
            raise ContextError("Pink Sheet header misses exact selected columns")
        positions = [period_position, *[header.index(name) for name in specs]]
        records = []
        physical_cap = max_rows // len(specs)
        last_data = sheet.max_row if data_end_row is None else min(data_end_row, sheet.max_row)
        if last_data - first_data > physical_cap:
            raise ContextLimitError("Pink Sheet worksheet extends beyond the selected row budget")
        for row in sheet.iter_rows(min_row=first_data + 1, max_row=last_data, values_only=True):
            values = [row[position] for position in positions]
            if all(value is None for value in values):
                continue
            if len(records) >= physical_cap:
                raise ContextLimitError("Pink Sheet selected observations exceed max_rows")
            records.append(dict(zip(selected, values, strict=True)))
        frame = pd.DataFrame(records, columns=selected)
    finally:
        book.close()
    return parse_pink_sheet(
        frame,
        period_column=period_name,
        column_specs=specs,
        received_at=received_at,
        revision_id=revision_id or hashlib.sha256(body).hexdigest(),
        period_format=period_format,
        published_at=published_at,
        availability_basis=availability_basis,
        max_rows=max_rows,
    )


def collect_pink_sheet(
    transport,
    store,
    *,
    snapshot_id,
    url=PINK_SHEET_URL,
    limits=None,
    source_metadata=None,
    offline=None,
    **parser_options,
):
    parsed = urlsplit(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname != "thedocs.worldbank.org"
        or not parsed.path.endswith(".xlsx")
        or parsed.query
        or parsed.fragment
    ):
        raise ContextError("Pink Sheet URL must be a direct official HTTPS XLSX file")
    options = dict(parser_options)
    options["column_specs"] = {
        name: asdict(spec) for name, spec in _columns(options.get("column_specs")).items()
    }
    return _collect_file(
        transport,
        store,
        source_id="world_bank",
        dataset="pink_sheet_monthly",
        url=url,
        snapshot_id=snapshot_id,
        options=options,
        parser=parse_pink_sheet_xlsx,
        limits=limits,
        source_metadata=source_metadata,
        offline=offline,
    )


def pink_sheet_values_asof(frame, decision_at):
    """Latest available economic month per factor; never fill before receipt/release."""
    visible = observations_asof(frame, decision_at)
    return (
        visible.sort_values("observation_date", kind="stable")
        .drop_duplicates(["entity_id", "metric", "source_id"], keep="last")
        .reset_index(drop=True)
    )
